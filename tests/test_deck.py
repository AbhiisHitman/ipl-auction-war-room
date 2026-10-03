"""The generated deck opens cleanly and every slide carries an action title."""
import pytest
from pptx import Presentation

from src.config import OUTPUTS_DIR, PROCESSED_DIR

pytestmark = pytest.mark.skipif(not (PROCESSED_DIR / "players.csv").exists(),
                                reason="run `python -m src.pipeline --fast` first")


def test_deck_builds_and_reopens(tmp_path):
    from src.build_deck import build
    path = build("SRH")
    prs = Presentation(path)
    assert len(prs.slides) == 14
    for i, slide in enumerate(prs.slides, start=1):
        if i == 1:
            continue  # title slide
        titles = [s for s in slide.shapes if s.name == "Title"]
        assert titles and len(titles[0].text_frame.text) > 20, f"slide {i} lacks an action title"
        assert not titles[0].text_frame.text.endswith("."), f"slide {i} title ends with a period"
    assert (OUTPUTS_DIR / "deck" / "auction_war_room_SRH.pptx").exists()
