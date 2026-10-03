"""Build the Excel deliverables with LIVE formulas (no hard-coded results).

outputs/excel/franchise_model.xlsx
    Assumptions  every input (value, unit, status, source) - named cells
    Teams        per-team inputs (brand value, ground, capacity, fee, salary)
    P&L          3 scenario columns, every cell a formula on the inputs
    Scenarios    probability-weighted revenue / profit, playoff-chance input
    Check        Python (src/economics.py) vs Excel for the selected team

outputs/excel/squad_optimiser_<TEAM>.xlsx  (one per franchise, mini-auction)
    Instructions how to run Excel Solver
    Model        candidates (own squad + top-60 pool), decision column, totals,
                 constraints and objective - Solver-ready (<= 200 variables)
    Python       the PuLP answer to the same problem, for comparison

    python -m src.build_excel
"""
from __future__ import annotations

import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .config import EXCEL_DIR, PROCESSED_DIR, assumptions, ensure_dirs, rules
from .economics import SCENARIOS, default_inputs, pnl
from .optimiser import optimise_simple, solver_candidates
from .squad import ROLES
from .teams import active_codes, full_name

NAVY = "1F3A5F"
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")  # yellow = editable input
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
CALC_FILL = PatternFill("solid", fgColor="EAF1F8")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def _name(wb: Workbook, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _header(ws, row: int, values: list[str]) -> None:
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.fill, c.font, c.border = HEAD_FILL, WHITE_BOLD, BOX
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(ws, widths: list[int]) -> None:
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


# ======================================================================
# Franchise model
# ======================================================================
GLOBAL_INPUTS = [
    # (name, label, path in assumptions/rules, unit)
    ("media_rights_total", "Central media rights, 2023-27 cycle", ("economics", "media_rights_total_cr"), "INR cr"),
    ("media_rights_seasons", "Seasons in media-rights cycle", ("economics", "media_rights_seasons"), "seasons"),
    ("share_media", "Franchise share of media rights", ("economics", "franchise_share_media_rights"), "share"),
    ("title_spons", "Title sponsorship per season", ("economics", "title_sponsorship_cr_per_season"), "INR cr"),
    ("other_central_spons", "Other central sponsorship per season", ("economics", "other_central_sponsorship_cr_per_season"), "INR cr"),
    ("share_spons", "Franchise share of central sponsorship", ("economics", "franchise_share_central_sponsorship"), "share"),
    ("n_teams", "Number of teams", ("economics", "n_teams"), "teams"),
    ("team_spons_base", "Team sponsorship, average brand", ("economics", "team_sponsorship_base_cr"), "INR cr"),
    ("merch_base", "Merchandise, average brand", ("economics", "merchandise_base_cr"), "INR cr"),
    ("home_matches", "Home matches", ("economics", "home_matches"), "matches"),
    ("occupancy", "Stadium occupancy", ("economics", "occupancy"), "share"),
    ("max_seats", "Max sellable seats", ("economics", "max_sellable_seats"), "seats"),
    ("comp_share", "Complimentary seat share", ("economics", "complimentary_share"), "share"),
    ("ticket_price", "Average ticket price", ("economics", "avg_ticket_price_inr"), "INR"),
    ("match_day_cost", "Match-day cost per home match", ("economics", "match_day_cost_cr_per_home_match"), "INR cr"),
    ("match_fee_lakh", "Match fee per player per match", ("rules", "match_fee_lakh"), "INR lakh"),
    ("match_fee_players", "Players paid a match fee per match", ("economics", "match_fee_players_per_match"), "players"),
    ("league_matches", "League matches per team", ("rules", "league_matches_per_team"), "matches"),
    ("support_staff", "Support staff", ("economics", "support_staff_cr"), "INR cr"),
    ("ops_admin", "Operations & admin", ("economics", "operations_admin_cr"), "INR cr"),
    ("marketing", "Marketing", ("economics", "marketing_cr"), "INR cr"),
    ("travel", "Travel", ("economics", "travel_cr"), "INR cr"),
    ("bcci_fee_share", "BCCI fee on team sponsorship + merch", ("economics", "bcci_fee_share_of_team_revenue"), "share"),
]


def _lookup(path: tuple[str, str]) -> dict:
    sec, key = path
    if sec == "rules":
        return rules()["common"][key]
    return assumptions()[sec][key]


def build_franchise_model(default_team: str = "CSK") -> str:
    a = assumptions()["economics"]
    homes = pd.read_csv(PROCESSED_DIR / "home_venues.csv").set_index("team")["venue"]
    players = pd.read_csv(PROCESSED_DIR / "players.csv")
    salary = players.groupby("current_team")["current_salary_lakh"].sum() / 100
    summary = pd.read_csv(PROCESSED_DIR / "optimiser_summary.csv")
    p_play = summary[summary["mode"] == "mini_2027"].set_index("team")["playoff_prob_before"]

    wb = Workbook()
    # ------------------------------------------------------------ Assumptions
    ws = wb.active
    ws.title = "Assumptions"
    ws["A1"] = "IPL Franchise Model - Assumptions"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = "Yellow cells are inputs: change any of them and the P&L, Scenarios and Check sheets update."
    ws["A3"] = "Selected team:"
    ws["A3"].font = Font(bold=True)
    ws["B3"] = default_team
    ws["B3"].fill = INPUT_FILL
    dv = DataValidation(type="list", formula1='"' + ",".join(active_codes()) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    _name(wb, "sel_team", "Assumptions!$B$3")
    _header(ws, 5, ["Input", "Value", "Unit", "Status", "Source / reason"])
    r = 6
    for name, label, path, unit in GLOBAL_INPUTS:
        item = _lookup(path)
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=item["value"])
        c.fill = INPUT_FILL
        ws.cell(row=r, column=3, value=unit)
        ws.cell(row=r, column=4, value=item.get("status", "verified"))
        ws.cell(row=r, column=5, value=item.get("source") or item.get("reason"))
        _name(wb, name, f"Assumptions!$B${r}")
        r += 1
    # Prize money and uplifts (dicts)
    dict_inputs = [
        ("prize_winner", "Prize money - winner", a["prize_money_cr"], "winner", "INR cr"),
        ("prize_runner", "Prize money - runner-up", a["prize_money_cr"], "runner_up", "INR cr"),
        ("prize_third", "Prize money - 3rd", a["prize_money_cr"], "third", "INR cr"),
        ("prize_fourth", "Prize money - 4th", a["prize_money_cr"], "fourth", "INR cr"),
        ("spons_up_playoffs", "Sponsorship uplift - made playoffs", a["sponsorship_uplift"], "made_playoffs", "share"),
        ("spons_up_title", "Sponsorship uplift - won title", a["sponsorship_uplift"], "won_title", "share"),
        ("merch_up_playoffs", "Merchandise uplift - made playoffs", a["merchandise_uplift"], "made_playoffs", "share"),
        ("merch_up_title", "Merchandise uplift - won title", a["merchandise_uplift"], "won_title", "share"),
    ]
    for name, label, item, key, unit in dict_inputs:
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=item["value"][key])
        c.fill = INPUT_FILL
        ws.cell(row=r, column=3, value=unit)
        ws.cell(row=r, column=4, value=item["status"])
        ws.cell(row=r, column=5, value=item.get("source") or item.get("reason"))
        _name(wb, name, f"Assumptions!$B${r}")
        r += 1
    for name, label, val, note in [
        ("po_games_playoffs", "Extra matches - made playoffs (non-winner)", 2.0, "assumption: 1-3 playoff games"),
        ("po_games_title", "Extra matches - won title", 2.5, "assumption: 2-3 playoff games"),
        ("title_given_playoffs", "P(title | playoffs)", assumptions()["strategy"]["title_prob_given_playoffs"]["value"],
         assumptions()["strategy"]["title_prob_given_playoffs"]["reason"]),
    ]:
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=val)
        c.fill = INPUT_FILL
        ws.cell(row=r, column=4, value="assumption")
        ws.cell(row=r, column=5, value=note)
        _name(wb, name, f"Assumptions!$B${r}")
        r += 1
    _widths(ws, [44, 14, 10, 12, 110])

    # ------------------------------------------------------------ Teams
    wt = wb.create_sheet("Teams")
    _header(wt, 1, ["Team", "Franchise", "Brand value (USD m)", "Home ground", "Capacity",
                    "Franchise fee (INR cr/yr)", "2026 squad salary (INR cr)", "Model playoff chance"])
    teams = active_codes()
    for i, t in enumerate(teams, start=2):
        venue = homes[t]
        vals = [t, full_name(t), a["brand_value_usd_m"]["value"][t], venue,
                a["stadium_capacity"]["value"].get(venue, 40000), a["franchise_fee_cr_per_season"]["value"][t],
                round(float(salary.get(t, 0)), 2), round(float(p_play.get(t, 0.5)), 3)]
        for j, v in enumerate(vals, start=1):
            c = wt.cell(row=i, column=j, value=v)
            if j >= 3 and j != 4:
                c.fill = INPUT_FILL
    last = 1 + len(teams)
    for col, nm in [("A", "t_code"), ("C", "t_brand"), ("E", "t_capacity"), ("F", "t_fee"),
                    ("G", "t_salary"), ("H", "t_playoff")]:
        _name(wb, nm, f"Teams!${col}$2:${col}${last}")
    wt.cell(row=last + 2, column=1, value="Sources: brand values - Houlihan Lokey IPL Valuation Study 2025; "
            "capacities - Wikipedia stadium pages; home ground derived from 2026 fixtures; "
            "salary from 2026 retention + auction records; playoff chance from src/impact_model.py.")
    _widths(wt, [8, 30, 18, 44, 11, 22, 24, 20])

    # ------------------------------------------------------------ P&L
    wp = wb.create_sheet("P&L")
    wp["A1"] = "Franchise P&L - INR crore per season"
    wp["A1"].font = Font(bold=True, size=14, color=NAVY)
    wp["A2"] = '="Team: "&sel_team'
    _header(wp, 4, ["Line", "Missed playoffs", "Made playoffs", "Won title", "How it is calculated"])
    # Helper rows (team inputs) first so formulas can reference them
    helpers = [
        ("Brand index (team / league avg)", "=INDEX(t_brand,MATCH(sel_team,t_code,0))/AVERAGE(t_brand)", "brand_index"),
        ("Stadium capacity", "=INDEX(t_capacity,MATCH(sel_team,t_code,0))", "capacity"),
        ("Player salaries (INR cr)", "=INDEX(t_salary,MATCH(sel_team,t_code,0))", "salary"),
        ("Franchise fee (INR cr)", "=INDEX(t_fee,MATCH(sel_team,t_code,0))", "fee"),
    ]
    r = 5
    for label, f, nm in helpers:
        wp.cell(row=r, column=1, value=label)
        c = wp.cell(row=r, column=2, value=f)
        c.fill = CALC_FILL
        _name(wb, nm, f"'P&L'!$B${r}")
        r += 1
    r += 1
    cols = {"missed_playoffs": "B", "made_playoffs": "C", "won_title": "D"}
    spons_up = {"B": "0", "C": "spons_up_playoffs", "D": "spons_up_title"}
    merch_up = {"B": "0", "C": "merch_up_playoffs", "D": "merch_up_title"}
    prize = {"B": "0", "C": "(prize_runner+prize_third+prize_fourth)/3", "D": "prize_winner"}
    po_games = {"B": "0", "C": "po_games_playoffs", "D": "po_games_title"}
    lines = [
        ("REVENUE", None, None),
        ("Central media rights share", lambda c: "=media_rights_total/media_rights_seasons*share_media/n_teams",
         "Cycle value / seasons x franchise share / 10 teams"),
        ("Central sponsorship share", lambda c: "=(title_spons+other_central_spons)*share_spons/n_teams",
         "(Title + other central sponsors) x franchise share / 10"),
        ("Team sponsorships", lambda c: f"=team_spons_base*brand_index*(1+{spons_up[c]})",
         "Base x brand index x (1 + performance uplift)"),
        ("Ticketing (home matches)", lambda c: "=home_matches*MIN(capacity,max_seats)*occupancy*(1-comp_share)*ticket_price/10^7",
         "Home matches x sellable seats x occupancy x paid share x price"),
        ("Merchandise", lambda c: f"=merch_base*brand_index*(1+{merch_up[c]})", "Base x brand index x (1 + uplift)"),
        ("Prize money", lambda c: f"={prize[c]}", "BCCI prize pot by finish"),
        ("Total revenue", "SUM_REV", "Sum of revenue lines"),
        ("COSTS", None, None),
        ("Player salaries", lambda c: "=salary", "2026 squad salary bill (Teams sheet)"),
        ("Match fees", lambda c: f"=match_fee_lakh*match_fee_players*(league_matches+{po_games[c]})/100",
         "Fee x 12 players x matches played"),
        ("Support staff", lambda c: "=support_staff", "Input"),
        ("Operations & admin", lambda c: "=ops_admin", "Input"),
        ("Marketing", lambda c: "=marketing", "Input"),
        ("Travel", lambda c: "=travel", "Input"),
        ("Match-day costs", lambda c: "=match_day_cost*home_matches", "Cost per home match x home matches"),
        ("BCCI fee on team revenue", "BCCI", "Share x (team sponsorships + merchandise)"),
        ("Total operating costs", "SUM_COST", "Sum of cost lines"),
        ("PROFIT", None, None),
        ("Operating profit", "OP", "Total revenue - total operating costs"),
        ("Operating margin", "MARGIN", "Operating profit / revenue"),
        ("Franchise fee instalment (assumption)", lambda c: "=fee", "LSG/GT 2021 bid / 10 years (assumption)"),
        ("Profit after franchise fee", "AFTER", "Operating profit - franchise fee"),
        ("Central (fixed) share of revenue", "MIX", "Central income / total revenue"),
    ]
    rowmap: dict[str, int] = {}
    for label, f, note in lines:
        wp.cell(row=r, column=1, value=label)
        if f is None:
            wp.cell(row=r, column=1).font = Font(bold=True, color=NAVY)
            r += 1
            continue
        rowmap[label] = r
        for c in cols.values():
            if callable(f):
                formula = f(c)
            elif f == "SUM_REV":
                formula = f"=SUM({c}{rowmap['Central media rights share']}:{c}{rowmap['Prize money']})"
            elif f == "BCCI":
                formula = f"=bcci_fee_share*({c}{rowmap['Team sponsorships']}+{c}{rowmap['Merchandise']})"
            elif f == "SUM_COST":
                formula = f"=SUM({c}{rowmap['Player salaries']}:{c}{rowmap['BCCI fee on team revenue']})"
            elif f == "OP":
                formula = f"={c}{rowmap['Total revenue']}-{c}{rowmap['Total operating costs']}"
            elif f == "MARGIN":
                formula = f"={c}{rowmap['Operating profit']}/{c}{rowmap['Total revenue']}"
            elif f == "AFTER":
                formula = f"={c}{rowmap['Operating profit']}-{c}{rowmap['Franchise fee instalment (assumption)']}"
            elif f == "MIX":
                formula = (f"=({c}{rowmap['Central media rights share']}+{c}{rowmap['Central sponsorship share']})"
                           f"/{c}{rowmap['Total revenue']}")
            cell = wp.cell(row=r, column=ord(c) - 64, value=formula)
            cell.number_format = "0.0%" if label in ("Operating margin", "Central (fixed) share of revenue") else "#,##0.0"
            cell.border = BOX
        wp.cell(row=r, column=5, value=note).font = Font(italic=True, color="666666")
        if label in ("Total revenue", "Total operating costs", "Operating profit", "Profit after franchise fee"):
            for j in range(1, 5):
                wp.cell(row=r, column=j).font = Font(bold=True)
        r += 1
    _widths(wp, [40, 16, 16, 16, 62])
    wp.freeze_panes = "B5"

    # ------------------------------------------------------------ Scenarios
    wsn = wb.create_sheet("Scenarios")
    wsn["A1"] = "What is performance worth? Probability-weighted P&L"
    wsn["A1"].font = Font(bold=True, size=14, color=NAVY)
    wsn["A3"], wsn["B3"] = "P(make playoffs)", "=INDEX(t_playoff,MATCH(sel_team,t_code,0))"
    wsn["B3"].fill = INPUT_FILL
    wsn["C3"] = "Defaults to the model's chance for the current squad; overwrite to test."
    wsn["A4"], wsn["B4"] = "P(title)", "=B3*title_given_playoffs"
    _header(wsn, 6, ["Scenario", "Probability", "Revenue", "Operating profit"])
    tr, op = rowmap["Total revenue"], rowmap["Operating profit"]
    sc = [("Missed playoffs", "=1-B3", "B"), ("Made playoffs (no title)", "=B3-B4", "C"), ("Won title", "=B4", "D")]
    for i, (lab, prob, c) in enumerate(sc, start=7):
        wsn.cell(row=i, column=1, value=lab)
        wsn.cell(row=i, column=2, value=prob).number_format = "0.0%"
        wsn.cell(row=i, column=3, value=f"='P&L'!{c}{tr}").number_format = "#,##0.0"
        wsn.cell(row=i, column=4, value=f"='P&L'!{c}{op}").number_format = "#,##0.0"
    wsn["A10"] = "Expected value"
    wsn["A10"].font = Font(bold=True)
    wsn["C10"] = "=SUMPRODUCT(B7:B9,C7:C9)"
    wsn["D10"] = "=SUMPRODUCT(B7:B9,D7:D9)"
    wsn["A12"] = "Revenue at stake: title vs missed playoffs"
    wsn["C12"] = "=C9-C7"
    wsn["A13"] = "...as % of missed-playoffs revenue"
    wsn["C13"] = "=C12/C7"
    wsn["C13"].number_format = "0.0%"
    for cell in ("C10", "D10", "C12"):
        wsn[cell].number_format = "#,##0.0"
    _widths(wsn, [42, 14, 14, 18])

    # ------------------------------------------------------------ Check
    wc = wb.create_sheet("Check")
    wc["A1"] = "Python vs Excel (base case, default inputs)"
    wc["A1"].font = Font(bold=True, size=14, color=NAVY)
    wc["A2"] = ("Python values come from src/economics.py with the same default inputs. If you change an input, "
                "Excel will differ from these stored values - that is expected.")
    _header(wc, 4, ["Team", "Line", "Scenario", "Python", "Excel (selected team)", "Difference"])
    rr = 5
    check_lines = ["Total revenue", "Total operating costs", "Operating profit"]
    for t in teams:
        inp = default_inputs(t, homes[t], salary_cr=float(salary.get(t, 0)))
        for s in SCENARIOS:
            pt = pnl(inp, s).set_index("line")["value_cr"]
            for line in check_lines:
                wc.cell(row=rr, column=1, value=t)
                wc.cell(row=rr, column=2, value=line)
                wc.cell(row=rr, column=3, value=s)
                wc.cell(row=rr, column=4, value=round(float(pt[line]), 4))
                col = cols[s]
                wc.cell(row=rr, column=5, value=f"=IF(A{rr}=sel_team,'P&L'!{col}{rowmap[line]},\"\")")
                wc.cell(row=rr, column=6, value=f'=IF(E{rr}="","",E{rr}-D{rr})')
                for j in (4, 5, 6):
                    wc.cell(row=rr, column=j).number_format = "#,##0.0000"
                rr += 1
    wc.conditional_formatting.add(f"F5:F{rr}", CellIsRule(operator="notBetween", formula=["-0.001", "0.001"],
                                                          fill=PatternFill("solid", fgColor="F8CBAD")))
    _widths(wc, [8, 24, 18, 14, 22, 14])

    path = EXCEL_DIR / "franchise_model.xlsx"
    wb.save(path)
    return str(path)


# ======================================================================
# Excel Solver workbook
# ======================================================================
def build_solver_workbook(team: str, players: pd.DataFrame, briefs: dict, pool_mask: pd.Series) -> dict:
    cand, lam = solver_candidates(players, team, briefs[team]["role_weights"], pool_mask)
    py = optimise_simple(cand, lam)
    R = rules()
    c = R["common"]
    purse = R["modes"]["mini_2027"]["purse_lakh"]["value"]
    need = assumptions()["strategy"]["role_min_counts"]["value"]

    wb = Workbook()
    wi = wb.active
    wi.title = "Instructions"
    steps = [
        f"Squad optimiser - {full_name(team)} - IPL 2027 mini-auction (reduced pool for Excel Solver)",
        "",
        "What this solves: pick the squad (keep own players or release them; buy from the pool) that maximises",
        "   total value - lambda x total cost, within the purse and the squad rules.",
        "   value = (Impact Score - replacement level) x role weight for this team's gaps (see Model sheet).",
        "",
        "How to run Solver (Data > Solver; enable the Solver add-in via Tools > Excel Add-ins if needed):",
        "  1. Set Objective: Model!$C$5   To: Max",
        f"  2. By Changing Variable Cells: Model!$O$12:$O${11 + len(cand)}",
        "  3. Subject to the Constraints:",
        f"       Model!$O$12:$O${11 + len(cand)} = binary",
        "       Model!$C$6 <= Model!$D$6      (purse)",
        "       Model!$C$7 >= Model!$D$7      (minimum spend)",
        "       Model!$C$8 >= Model!$D$8      (squad minimum)",
        "       Model!$C$9 <= Model!$D$9      (squad maximum)",
        "       Model!$C$10 <= Model!$D$10    (overseas cap)",
        "       Model!$G$6:$G$10 >= Model!$H$6:$H$10 and Model!$J$6:$J$10 >= Model!$K$6:$K$10  (role minimums)",
        "  4. Solving method: Simplex LP. Click Solve.",
        "",
        "The decision column starts at the PuLP (Python) answer. Re-solving in Excel should give the same or",
        "a near-identical squad and objective - compare with the 'Python' sheet.",
        f"Variables: {len(cand)} (Excel Solver limit 200). Roles with too few candidates have their minimum set to the number available.",
    ]
    for i, s in enumerate(steps, start=1):
        wi.cell(row=i, column=1, value=s)
    wi["A1"].font = Font(bold=True, size=14, color=NAVY)
    wi.column_dimensions["A"].width = 120

    wm = wb.create_sheet("Model")
    wm["A1"] = f"{full_name(team)} - Solver model"
    wm["A1"].font = Font(bold=True, size=14, color=NAVY)
    first, last = 12, 11 + len(cand)
    rng = lambda col: f"${col}${first}:${col}${last}"  # noqa: E731
    wm["A3"], wm["B3"] = "lambda (value per lakh)", round(lam, 6)
    wm["B3"].fill = INPUT_FILL
    _header(wm, 4, ["Totals", "", "Value", "Limit"])
    wm["A5"], wm["C5"] = "OBJECTIVE: value - lambda x cost", f"=SUMPRODUCT({rng('M')},{rng('O')})-B3*C6"
    wm["A6"], wm["C6"], wm["D6"] = "Total cost (lakh)", f"=SUMPRODUCT({rng('N')},{rng('O')})", purse
    wm["A7"], wm["C7"], wm["D7"] = "Minimum spend (lakh)", "=C6", c["min_spend_pct"]["value"] * purse
    wm["A8"], wm["C8"], wm["D8"] = "Squad size (min)", f"=SUM({rng('O')})", c["squad_min"]["value"]
    wm["A9"], wm["C9"], wm["D9"] = "Squad size (max)", "=C8", c["squad_max"]["value"]
    wm["A10"], wm["C10"], wm["D10"] = "Overseas players", f"=SUMPRODUCT({rng('D')},{rng('O')})", c["overseas_max"]["value"]
    for cell in ("C5", "C6", "C8", "C10"):
        wm[cell].fill = CALC_FILL
    for cell in ("D6", "D7", "D8", "D9", "D10"):
        wm[cell].fill = INPUT_FILL
    # Role constraint block (two columns of 5)
    wm["F4"], wm["G4"], wm["H4"] = "Role", "In squad", "Min"
    wm["I4"], wm["J4"], wm["K4"] = "Role", "In squad", "Min"
    role_cols = {}
    head = ["Player", "Owner", "Role (primary)", "Overseas", "Impact", "Role weight"] + \
           [f"R:{r}" for r in ROLES[:6]]
    # Columns: A Player, B Owner, C Primary, D Overseas(1/0), E Impact, F RoleW,
    # then role flags in P..Y, M value, N cost, O pick.
    _header(wm, 11, ["Player", "Owner", "Primary role", "Overseas", "Impact", "Role weight",
                     "Salary/expected cost (cr)", "", "", "", "", "Value", "Value", "Cost (lakh)", "PICK (0/1)"]
            + [r for r in ROLES])
    for k, r_ in enumerate(ROLES):
        role_cols[r_] = get_column_letter(16 + k)
    for i, row in cand.iterrows():
        rr = first + i
        wm.cell(row=rr, column=1, value=row["display_name"])
        wm.cell(row=rr, column=2, value=row["owner"])
        wm.cell(row=rr, column=3, value=row.get("primary_role"))
        wm.cell(row=rr, column=4, value=int(row["is_overseas"]))
        wm.cell(row=rr, column=5, value=None if pd.isna(row["impact"]) else round(float(row["impact"]), 3))
        wm.cell(row=rr, column=6, value=round(float(row["role_w"]), 3))
        wm.cell(row=rr, column=7, value=round(float(row["cost_lakh"]) / 100, 2))
        wm.cell(row=rr, column=13, value=float(row["value"]))
        wm.cell(row=rr, column=14, value=float(row["cost_lakh"]))
        pick = wm.cell(row=rr, column=15, value=int(py[i]))
        pick.fill = INPUT_FILL
        for r_ in ROLES:
            wm[f"{role_cols[r_]}{rr}"] = int(r_ in row["roles"])
    for k, r_ in enumerate(ROLES):
        holders = int(sum(r_ in rl for rl in cand["roles"]))
        minimum = min(int(need.get(r_, 0)), holders)
        col_block = ("F", "G", "H") if k < 5 else ("I", "J", "K")
        rr = 5 + (k % 5) + 1
        wm[f"{col_block[0]}{rr}"] = r_
        wm[f"{col_block[1]}{rr}"] = f"=SUMPRODUCT({rng(role_cols[r_])},{rng('O')})"
        wm[f"{col_block[2]}{rr}"] = minimum
        wm[f"{col_block[1]}{rr}"].fill = CALC_FILL
        wm[f"{col_block[2]}{rr}"].fill = INPUT_FILL
    _widths(wm, [26, 8, 18, 9, 9, 11, 14, 2, 20, 10, 6, 2, 9, 11, 11] + [9] * len(ROLES))
    wm.freeze_panes = "B12"

    wpy = wb.create_sheet("Python")
    wpy["A1"] = "PuLP solution of the same reduced problem"
    wpy["A1"].font = Font(bold=True, size=14, color=NAVY)
    sel = cand[py.astype(bool).values]
    obj = float((sel["value"]).sum() - lam * sel["cost_lakh"].sum())
    wpy["A2"] = f"Objective {obj:.3f} | cost Rs {sel['cost_lakh'].sum() / 100:.2f} cr | squad {len(sel)} | overseas {int(sel['is_overseas'].sum())}"
    _header(wpy, 4, ["Player", "Owner", "Primary role", "Impact", "Cost (cr)", "Decision"])
    for k, (_, row) in enumerate(cand.iterrows(), start=5):
        dec = ("Keep" if row["owner"] == "own" else "Buy") if py[row.name] else ("Release" if row["owner"] == "own" else "-")
        for j, v in enumerate([row["display_name"], row["owner"], row.get("primary_role"),
                               None if pd.isna(row["impact"]) else round(float(row["impact"]), 2),
                               round(float(row["cost_lakh"]) / 100, 2), dec], start=1):
            wpy.cell(row=k, column=j, value=v)
    _widths(wpy, [26, 8, 18, 9, 10, 10])

    path = EXCEL_DIR / f"squad_optimiser_{team}.xlsx"
    wb.save(path)
    return {"team": team, "path": str(path), "objective": obj, "variables": len(cand)}


def main() -> None:
    ensure_dirs()
    print(build_franchise_model())
    players = pd.read_csv(PROCESSED_DIR / "players.csv")
    players["roles"] = players["roles"].fillna("").apply(lambda s: [r for r in s.split(";") if r])
    briefs = json.loads((PROCESSED_DIR / "team_briefs.json").read_text())
    alive = ~players["retired"] & players["overseas"].notna()
    recent = players["contract_season"].fillna(0).ge(2025) | players["last_season"].fillna(0).ge(2025)
    mask = alive & recent & (players["current_team"].isna() | players["projected_release"])
    for t in active_codes():
        print(build_solver_workbook(t, players, briefs, mask))


if __name__ == "__main__":
    main()
